"""
export_workflow_docx.py
Xuất tài liệu kết nối 2 báo cáo SLA thành DOCX — lồng cả 2 dashboard.

Usage:
    python3 export_workflow_docx.py
Output:
    exports/sla-workflow.docx
"""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE       = Path(__file__).parent
EXPORTS    = BASE / 'exports'
IMG_BOD    = BASE / 'order-sla/exports/dashboard.png'
IMG_OPS    = BASE / 'ops-sla-drilldown/exports/dashboard.png'
EXPORTS.mkdir(exist_ok=True)
OUT        = EXPORTS / 'sla-workflow.docx'

for img in [IMG_BOD, IMG_OPS]:
    assert img.exists(), f'Missing: {img} — chạy export_dashboard.py trong folder tương ứng trước.'

doc = Document()

for section in doc.sections:
    section.top_margin    = Inches(0.79)
    section.bottom_margin = Inches(0.79)
    section.left_margin   = Inches(0.87)
    section.right_margin  = Inches(0.87)


# ── Helpers ───────────────────────────────────────────────────────────────────
def add_rule(doc, color='CCCCCC'):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), color)
    pBdr.append(bottom)
    pPr.append(pBdr)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(2)


def add_para(doc, text, size=10, bold=False, color=None, italic=False, space_after=6):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size    = Pt(size)
    r.bold         = bold
    r.italic       = italic
    if color:
        r.font.color.rgb = RGBColor(*color)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def add_img_centered(doc, img_path, width=Inches(6.2)):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(4)
    p.add_run().add_picture(str(img_path), width=width)


def add_table(doc, headers, rows, col_widths=None):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.style = 'Table Grid'
    for i, (cell, text) in enumerate(zip(tbl.rows[0].cells, headers)):
        cell.text = text
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(9.5)
        if col_widths:
            cell.width = col_widths[i]
    for row_data in rows:
        row = tbl.add_row()
        for cell, val in zip(row.cells, row_data):
            cell.text = val
            cell.paragraphs[0].runs[0].font.size = Pt(9.5)
    return tbl


def add_flow_box(doc, lines, bg_color=None):
    """Thêm box text kiểu flowchart dùng border paragraph."""
    for line in lines:
        p = doc.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        # all-side border
        pBdr = OxmlElement('w:pBdr')
        for side in ['top', 'left', 'bottom', 'right']:
            el = OxmlElement(f'w:{side}')
            el.set(qn('w:val'), 'single')
            el.set(qn('w:sz'), '4')
            el.set(qn('w:space'), '4')
            el.set(qn('w:color'), 'AAAAAA')
            pBdr.append(el)
        pPr.append(pBdr)
        r = p.add_run(line)
        r.font.size = Pt(9.5)
        r.font.name = 'Courier New'
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after  = Pt(0)
        p.paragraph_format.left_indent  = Inches(0.2)


# ─────────────────────────────────────────────────────────────────────────────
# PAGE 1 — Title + Tổng quan + Luồng
# ─────────────────────────────────────────────────────────────────────────────
title = doc.add_heading('Luồng Báo Cáo SLA — Từ BOD đến Ops', level=1)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

add_para(doc,
    'Tài liệu này giải thích 2 báo cáo SLA kết nối với nhau thế nào: '
    'ai đọc gì, khi nào escalate, và ops team báo cáo lại ra sao.',
    size=10, color=(0x55, 0x55, 0x55), italic=True)

add_rule(doc)

# ── Tổng quan 2 tầng ─────────────────────────────────────────────────────────
h = doc.add_heading('Hệ Thống 2 Tầng', level=2)

add_table(doc,
    headers=['Tầng', 'Báo cáo', 'Audience', 'Câu hỏi', 'Tần suất'],
    rows=[
        ('L1 — BOD', 'Order SLA',
         'BOD, leadership',
         'Tháng này đạt bao nhiêu? Trend đi đâu?',
         'Hàng tháng'),
        ('L2 — Ops', 'Ops SLA Drilldown',
         'Ops team',
         'Xưởng nào? Sản phẩm nào? Nguyên nhân gì?',
         'Khi được escalate'),
    ]
)

doc.add_paragraph()
add_para(doc,
    'Nguyên tắc: BOD không drill down. Ops không báo cáo aggregate. '
    'Hai tầng bổ sung nhau — không thay thế nhau.',
    bold=True, size=10)

add_rule(doc)

# ── Luồng hành động ──────────────────────────────────────────────────────────
doc.add_heading('Luồng Đọc và Hành Động', level=2)

flow_lines = [
    '  BOD đọc Order SLA Dashboard (đầu tháng)                              ',
    '    Phần 1 — Scorecard: Platform XX% ▲/▼  ·  Standard XX% ▲/▼          ',
    '    Phần 2 — Trend: đang tốt lên hay xấu đi?                           ',
    '                                                                        ',
    '         Tình hình ổn?                   Cần escalate?                 ',
    '              │                                │                        ',
    '              ▼                                ▼                        ',
    '    Không cần action            Gửi cho Ops: "Tháng X: [nhóm]          ',
    '    Ghi nhận, theo dõi           drop [Y]pp → cần drill-down"          ',
    '    tháng tới                               │                           ',
    '                                            ▼                           ',
    '                           Ops đọc Ops SLA Drilldown                   ',
    '                             Bước 1 — Xưởng nào chậm?                  ',
    '                             Bước 2 — Sản phẩm nào ưu tiên?            ',
    '                             Bước 3 — Carrier hay sản xuất?             ',
    '                                            │                           ',
    '                                            ▼                           ',
    '                           Ops báo cáo lại BOD (template)              ',
    '                             — Vấn đề ở đâu                            ',
    '                             — Đang làm gì / timeline                  ',
    '                             — Cần BOD support gì không                ',
]
add_flow_box(doc, flow_lines)

add_rule(doc)

# ── Khi nào escalate ─────────────────────────────────────────────────────────
doc.add_heading('BOD — Khi Nào Escalate?', level=2)

add_para(doc, 'Escalate ngay (tháng này):', bold=True, size=10, space_after=2)
add_table(doc,
    headers=['Điều kiện', 'Ví dụ'],
    rows=[
        ('Bất kỳ nhóm nào < 85%',       'Platform 82% tháng này'),
        ('Drop > 10pp vs tháng trước',   'Standard 96% → 85%'),
        ('Cả 2 nhóm cùng drop 1 tháng', 'Pattern Dec 2025 lặp lại'),
    ]
)

doc.add_paragraph()
add_para(doc, 'Theo dõi sát (tháng tới):', bold=True, size=10, space_after=2)
add_table(doc,
    headers=['Điều kiện', 'Ví dụ'],
    rows=[
        ('Drop 5–10pp vs tháng trước',         'Platform 96% → 91%'),
        ('Trending down 2 tháng liên tiếp',    'Standard giảm liên tục'),
        ('Volume tăng mạnh sắp đến',           'Tháng cao điểm, Q4'),
    ]
)

doc.add_paragraph()
add_para(doc, 'Không cần action: Cả 2 nhóm ≥ 90% và delta ≥ −3pp so với tháng trước.',
         size=10, color=(0x27, 0xae, 0x60))

add_rule(doc)

# ── Lịch vận hành ────────────────────────────────────────────────────────────
doc.add_heading('Lịch Vận Hành Hàng Tháng', level=2)

add_table(doc,
    headers=['Thời điểm', 'Ai', 'Hành động'],
    rows=[
        ('Ngày 3–5', 'Data team',   'Refresh CSV, re-export 2 dashboard + 2 DOCX'),
        ('Ngày 5–7', 'BOD',         'Đọc Order SLA — quyết định escalate hay không'),
        ('Ngày 7–10 (nếu escalate)', 'Ops', 'Đọc Ops Drilldown, xác định root cause'),
        ('Ngày 10–12 (nếu escalate)', 'Ops → BOD', 'Gửi Ops Update theo template'),
        ('Cuối tháng', 'Ops',       'Check progress — đúng hướng chưa?'),
    ]
)

# ─────────────────────────────────────────────────────────────────────────────
# PAGE BREAK → BOD Dashboard
# ─────────────────────────────────────────────────────────────────────────────
doc.add_page_break()

h = doc.add_heading('L1 — BOD Dashboard: Order SLA', level=1)
h.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_para(doc,
    'Đọc hàng tháng. Trả lời: Tháng này đạt bao nhiêu? Trend đang đi đâu?',
    size=10, italic=True, color=(0x55, 0x55, 0x55))

add_img_centered(doc, IMG_BOD, width=Inches(6.2))

add_rule(doc)
doc.add_heading('Cách đọc nhanh (≤ 30 giây)', level=2)

steps_bod = [
    ('1', 'Xem status box (header xanh/vàng)',
     'Nếu xanh → không cần escalate. Nếu vàng → đọc tiếp.'),
    ('2', 'Phần 1 — Scorecard',
     'Hai con số lớn + delta. Đây là số cần quyết định.'),
    ('3', 'Phần 2 — Trend',
     'Dec 2025 là outlier (holiday). Recovery đang diễn ra — đây là bối cảnh bình thường.'),
    ('4', 'Escalate nếu cần',
     'Gửi screenshot + "Tháng X: [nhóm] drop [Y]pp" cho ops team.'),
]
tbl = doc.add_table(rows=1, cols=3)
tbl.style = 'Table Grid'
for cell, text in zip(tbl.rows[0].cells, ['Bước', 'Nhìn vào', 'Làm gì']):
    cell.text = text
    cell.paragraphs[0].runs[0].bold = True
    cell.paragraphs[0].runs[0].font.size = Pt(9.5)
for step, look, action in steps_bod:
    row = tbl.add_row()
    for cell, val in zip(row.cells, [step, look, action]):
        cell.text = val
        cell.paragraphs[0].runs[0].font.size = Pt(9.5)

# ─────────────────────────────────────────────────────────────────────────────
# PAGE BREAK → Ops Dashboard
# ─────────────────────────────────────────────────────────────────────────────
doc.add_page_break()

h = doc.add_heading('L2 — Ops Dashboard: SLA Drilldown', level=1)
h.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_para(doc,
    'Đọc khi được escalate. Trả lời: Xưởng nào? Sản phẩm nào? Carrier hay sản xuất?',
    size=10, italic=True, color=(0x55, 0x55, 0x55))

add_img_centered(doc, IMG_OPS, width=Inches(6.2))

add_rule(doc)
doc.add_heading('Cách đọc theo bước', level=2)

steps_ops = [
    ('Bước 1', 'Xưởng nào chậm?',
     'So sánh CA vs TX1. CA = HB1+HB2 cùng ops lead. TX1 = chuẩn so sánh. '
     'Xác định xưởng cần focus trước.'),
    ('Bước 2', 'Sản phẩm nào ưu tiên?',
     'Trong xưởng đó, sản phẩm nào pass rate thấp nhất × volume lớn nhất? '
     'Đỏ = cần xử lý, xanh = ổn.'),
    ('Bước 3', 'Nguyên nhân gì?',
     'Bar stacked: xanh (carrier) vs cam (sản xuất). Phần nào cao hơn đường SLA = bottleneck. '
     'Label "↑ Carrier" hoặc "↑ Sản xuất" hiển thị sẵn.'),
]
tbl2 = doc.add_table(rows=1, cols=3)
tbl2.style = 'Table Grid'
for cell, text in zip(tbl2.rows[0].cells, ['', 'Câu hỏi', 'Làm gì với thông tin đó']):
    cell.text = text
    cell.paragraphs[0].runs[0].bold = True
    cell.paragraphs[0].runs[0].font.size = Pt(9.5)
for step, question, action in steps_ops:
    row = tbl2.add_row()
    for cell, val in zip(row.cells, [step, question, action]):
        cell.text = val
        cell.paragraphs[0].runs[0].font.size = Pt(9.5)

add_rule(doc)

# ── Template báo cáo lại ─────────────────────────────────────────────────────
doc.add_heading('Template Ops Báo Cáo Lại BOD', level=2)

template_lines = [
    'Ops SLA Update — Tháng [X/XXXX]                                        ',
    '                                                                        ',
    'Vấn đề xác định:                                                        ',
    '  - [Facility]: [Sản phẩm] — pass rate [X]%, nguyên nhân [carrier/sx] ',
    '  - (thêm nếu có nhiều hơn 1 vấn đề)                                   ',
    '                                                                        ',
    'Đang xử lý:                                                             ',
    '  - [Hành động cụ thể]                                                  ',
    '  - [Timeline dự kiến hoàn thành]                                       ',
    '                                                                        ',
    'Kỳ vọng:                                                                ',
    '  - Tháng [X+1]: pass rate [facility/sản phẩm] cải thiện lên ~[X]%    ',
    '  - Nếu không cải thiện: cần leo thang về [carrier / ops process]      ',
    '                                                                        ',
    'Cần BOD support:                                                        ',
    '  - [Nếu có: ngân sách đổi carrier, tăng headcount, v.v.]              ',
    '  - "Không có gì" nếu ops tự xử lý được                                ',
]
add_flow_box(doc, template_lines)

# ─────────────────────────────────────────────────────────────────────────────
# Save
# ─────────────────────────────────────────────────────────────────────────────
doc.save(OUT)
print(f'Saved: {OUT}')
