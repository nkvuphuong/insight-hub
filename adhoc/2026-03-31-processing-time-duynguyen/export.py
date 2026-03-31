"""
BI-792: Export processing time report to Excel
Seller: duynguyen02169@gmail.com — Q1 2026
2 sheets: Summary (matrix) + Detail (raw list)
"""

import csv
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUTPUT_DIR = Path(__file__).parent / "output"

# Tìm file CSV mới nhất theo prefix (hỗ trợ cả tên cũ lẫn timestamped)
def latest_csv(prefix):
    files = sorted(OUTPUT_DIR.glob(f"{prefix}*.csv"), key=lambda f: f.stat().st_mtime, reverse=True)
    if not files:
        raise FileNotFoundError(f"Không tìm thấy CSV với prefix '{prefix}' trong {OUTPUT_DIR}")
    return files[0]

SUMMARY_CSV = latest_csv("query_summary")
DETAIL_CSV  = latest_csv("query_detail")
EXCEL_OUT   = OUTPUT_DIR / "BI-792_processing-time_duynguyen_Q1-2026.xlsx"

# ── Styles ────────────────────────────────────────────────────────────────────
HEADER_FILL   = PatternFill("solid", fgColor="1F4E79")
SUBHEAD_FILL  = PatternFill("solid", fgColor="2E75B6")
ALT_ROW_FILL  = PatternFill("solid", fgColor="EBF3FB")
HEADER_FONT   = Font(bold=True, color="FFFFFF", size=11)
SUBHEAD_FONT  = Font(bold=True, color="FFFFFF", size=10)
BOLD_FONT     = Font(bold=True, size=10)
NORMAL_FONT   = Font(size=10)
CENTER        = Alignment(horizontal="center", vertical="center")
LEFT          = Alignment(horizontal="left",   vertical="center")
thin_side     = Side(style="thin", color="BDD7EE")
THIN_BORDER   = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)


def set_header(cell, value):
    cell.value = value
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = CENTER
    cell.border = THIN_BORDER


def set_subheader(cell, value):
    cell.value = value
    cell.font = SUBHEAD_FONT
    cell.fill = SUBHEAD_FILL
    cell.alignment = CENTER
    cell.border = THIN_BORDER


def set_cell(cell, value, bold=False, center=False, alt=False):
    cell.value = value
    cell.font  = BOLD_FONT if bold else NORMAL_FONT
    cell.alignment = CENTER if center else LEFT
    cell.border = THIN_BORDER
    if alt:
        cell.fill = ALT_ROW_FILL


# ── Load CSVs ─────────────────────────────────────────────────────────────────
def load_csv(path):
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))

summary_rows = load_csv(SUMMARY_CSV)
detail_rows  = load_csv(DETAIL_CSV)

wb = openpyxl.Workbook()

# ══════════════════════════════════════════════════════════════════════════════
# SHEET 1: SUMMARY — Pivot (product_type × paid_month, value = avg_processing_day)
# ══════════════════════════════════════════════════════════════════════════════
ws1 = wb.active
ws1.title = "Summary"

# Build pivot structure from CSV
months        = sorted({r["paid_month"] for r in summary_rows})
product_types = sorted({r["product_type"] for r in summary_rows})
pivot = {(r["product_type"], r["paid_month"]): float(r["avg_processing_day"]) for r in summary_rows}

num_cols = 1 + len(months)  # product_type col + one col per month

# Title
ws1.merge_cells(f"A1:{get_column_letter(num_cols)}1")
title_cell = ws1["A1"]
title_cell.value = "Avg Processing Time (days) — duynguyen02169@gmail.com | Q1 2026"
title_cell.font  = Font(bold=True, size=13, color="1F4E79")
title_cell.alignment = CENTER
ws1.row_dimensions[1].height = 28

ws1.merge_cells(f"A2:{get_column_letter(num_cols)}2")
ws1["A2"].value = "Nguồn: fact_order_fulfillment_performance_by_product | avg theo product type × paid month"
ws1["A2"].font = Font(italic=True, size=9, color="595959")
ws1["A2"].alignment = CENTER
ws1.row_dimensions[2].height = 16

# Header row: "Product Type" | month1 | month2 | month3 ...
set_header(ws1.cell(row=3, column=1), "Product Type")
for col, m in enumerate(months, 2):
    set_header(ws1.cell(row=3, column=col), m)
ws1.row_dimensions[3].height = 22

# Data rows
for i, pt in enumerate(product_types):
    r = i + 4
    alt = (i % 2 == 1)
    set_cell(ws1.cell(r, 1), pt, bold=True, alt=alt)
    for col, m in enumerate(months, 2):
        val = pivot.get((pt, m))
        set_cell(ws1.cell(r, col), val, center=True, alt=alt)
        if val is not None:
            ws1.cell(r, col).number_format = "0.00"

# Column widths
ws1.column_dimensions["A"].width = 18
for col in range(2, num_cols + 1):
    ws1.column_dimensions[get_column_letter(col)].width = 16

# Freeze header
ws1.freeze_panes = "B4"

# ══════════════════════════════════════════════════════════════════════════════
# SHEET 2: DETAIL
# ══════════════════════════════════════════════════════════════════════════════
ws2 = wb.create_sheet("Detail")

ws2.merge_cells("A1:G1")
ws2["A1"].value = "Chi tiết đơn hàng — duynguyen02169@gmail.com | Q1 2026 (tháng 3)"
ws2["A1"].font  = Font(bold=True, size=13, color="1F4E79")
ws2["A1"].alignment = CENTER
ws2.row_dimensions[1].height = 28

detail_headers = [
    "Order ID", "Paid Date", "Product Type",
    "Workshop", "Processing (days)", "First Scan (days)", "Shipping (days)"
]
for col, h in enumerate(detail_headers, 1):
    set_header(ws2.cell(row=2, column=col), h)
ws2.row_dimensions[2].height = 22

for i, row in enumerate(detail_rows):
    r = i + 3
    alt = (i % 2 == 1)

    def flt(v):
        return float(v) if v else None

    set_cell(ws2.cell(r, 1), int(row["order_id"]),         center=True, alt=alt)
    set_cell(ws2.cell(r, 2), row["paid_date"],              center=True, alt=alt)
    set_cell(ws2.cell(r, 3), row["product_type"],           alt=alt)
    set_cell(ws2.cell(r, 4), row["workshop_code"],          center=True, alt=alt)
    set_cell(ws2.cell(r, 5), flt(row["processing_day"]),    center=True, alt=alt)
    set_cell(ws2.cell(r, 6), flt(row["first_scan_day"]),    center=True, alt=alt)
    set_cell(ws2.cell(r, 7), flt(row["shipping_day"]),      center=True, alt=alt)
    for col in [5, 6, 7]:
        c = ws2.cell(r, col)
        if c.value is not None:
            c.number_format = "0.00"

col_widths2 = [14, 14, 16, 12, 20, 20, 18]
for col, w in enumerate(col_widths2, 1):
    ws2.column_dimensions[get_column_letter(col)].width = w

ws2.freeze_panes = "A3"

# ── Save ──────────────────────────────────────────────────────────────────────
wb.save(EXCEL_OUT)
print(f"Saved: {EXCEL_OUT}")
